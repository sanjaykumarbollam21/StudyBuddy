import io
from typing import List
import docx
from app.documents.extractors.base import BaseExtractor, ExtractedDocument, ExtractedPage, ExtractedSection

class DOCXExtractor(BaseExtractor):
    def extract(self, content: bytes, filename: str = "") -> ExtractedDocument:
        try:
            doc = docx.Document(io.BytesIO(content))
        except Exception as e:
            raise ValueError(f"Corrupted or invalid DOCX document: {str(e)}")

        paragraphs_text: List[str] = []
        sections: List[ExtractedSection] = []
        char_cursor = 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name.lower() if para.style else ""

            # Check if paragraph is a heading
            if "heading" in style_name or "title" in style_name:
                level = 1
                if "2" in style_name:
                    level = 2
                elif "3" in style_name:
                    level = 3
                sections.append(
                    ExtractedSection(
                        title=text,
                        level=level,
                        character_start=char_cursor,
                    )
                )

            paragraphs_text.append(text)
            char_cursor += len(text) + 2

        # Extract table text
        for table_idx, table in enumerate(doc.tables):
            table_rows: List[str] = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(" | ".join(row_cells))
            if table_rows:
                table_str = f"\n[Table {table_idx + 1}]\n" + "\n".join(table_rows)
                paragraphs_text.append(table_str)

        full_text = "\n\n".join(paragraphs_text)
        # Approximate page count (standard page ~ 250 words / 1500 chars)
        estimated_pages = max(1, len(full_text) // 1500 + (1 if len(full_text) % 1500 else 0))

        pages = [
            ExtractedPage(page_number=i + 1, text=full_text[i * 1500 : (i + 1) * 1500])
            for i in range(estimated_pages)
        ]

        return ExtractedDocument(
            text=full_text,
            pages=pages,
            sections=sections,
            metadata={"paragraph_count": len(paragraphs_text), "extractor": "python-docx"},
            needs_ocr=False,
            page_count=estimated_pages,
        )
